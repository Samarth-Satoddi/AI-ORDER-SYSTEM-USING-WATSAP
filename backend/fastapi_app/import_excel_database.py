import argparse
import os
import re
import sys
from collections.abc import Iterable
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import OperationalError
from sqlalchemy.orm import Session

from fastapi_app.db.base import Base
from fastapi_app.db.session import engine
from fastapi_app.db.session import SessionLocal
from fastapi_app.models.hotel import Hotel
from fastapi_app.models.menu_item import MenuItem

MENU_REQUIRED_COLUMNS = {"Hotel ID", "Hotel Name", "Item Name", "Price (₹)"}
HOTEL_REQUIRED_COLUMNS = {"Hotel ID", "Hotel Name"}


@dataclass(frozen=True)
class ImportResult:
    hotels_created: int = 0
    hotels_updated: int = 0
    menu_items_created: int = 0
    menu_items_updated: int = 0
    menu_items_deactivated: int = 0


@dataclass(frozen=True)
class ExcelSummary:
    hotels: int
    menu_items: int
    menu_hotel_ids: int


def slugify(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", value.strip().lower())
    return slug.strip("-") or "hotel"


def normalized_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def decimal_price(value: Any) -> Decimal:
    if value is None or value == "":
        raise ValueError("price is required")
    try:
        return Decimal(str(value)).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"invalid price {value!r}") from exc


def yes_no(value: Any, default: bool = True) -> bool:
    if value is None or value == "":
        return default
    return str(value).strip().lower() in {"yes", "y", "true", "1", "available", "active"}


def is_summary_row(value: str) -> bool:
    return value.strip().lower() in {"total", "totals", "grand total"}


def row_dicts(sheet: Any, required_columns: set[str]) -> Iterable[dict[str, Any]]:
    header: list[str] | None = None

    for row in sheet.iter_rows(values_only=True):
        values = [normalized_text(value) for value in row]
        if header is None:
            if required_columns.issubset(set(values)):
                header = values
            continue

        record = {column: row[index] for index, column in enumerate(header) if column}
        if any(value not in (None, "") for value in record.values()):
            yield record


def find_sheet(workbook: Any, required_columns: set[str]) -> Any:
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            values = {normalized_text(value) for value in row}
            if required_columns.issubset(values):
                return sheet
    raise ValueError(f"Could not find a sheet with columns: {', '.join(sorted(required_columns))}")


def inspect_excel(path: Path) -> ExcelSummary:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        hotel_sheet = find_sheet(workbook, HOTEL_REQUIRED_COLUMNS)
        menu_sheet = find_sheet(workbook, MENU_REQUIRED_COLUMNS)

        hotels = [
            record
            for record in row_dicts(hotel_sheet, HOTEL_REQUIRED_COLUMNS)
            if not is_summary_row(normalized_text(record.get("Hotel ID")))
        ]
        menu_items = [
            record
            for record in row_dicts(menu_sheet, MENU_REQUIRED_COLUMNS)
            if not is_summary_row(normalized_text(record.get("Hotel ID")))
        ]
        menu_hotel_ids = {normalized_text(record.get("Hotel ID")) for record in menu_items}
        menu_hotel_ids.discard("")
        return ExcelSummary(hotels=len(hotels), menu_items=len(menu_items), menu_hotel_ids=len(menu_hotel_ids))
    finally:
        workbook.close()


def import_excel(path: Path, db: Session | None = None, deactivate_missing: bool = False) -> ImportResult:
    owns_session = db is None
    db = db or SessionLocal()
    workbook = load_workbook(path, read_only=True, data_only=True)

    hotels_created = hotels_updated = 0
    menu_items_created = menu_items_updated = menu_items_deactivated = 0
    imported_menu_keys: set[tuple[int, str]] = set()
    imported_hotel_ids: set[int] = set()

    try:
        hotel_sheet = find_sheet(workbook, HOTEL_REQUIRED_COLUMNS)
        menu_sheet = find_sheet(workbook, MENU_REQUIRED_COLUMNS)
        hotels_by_external_id: dict[str, Hotel] = {}

        for record in row_dicts(hotel_sheet, HOTEL_REQUIRED_COLUMNS):
            external_id = normalized_text(record.get("Hotel ID"))
            name = normalized_text(record.get("Hotel Name"))
            if not external_id or not name or is_summary_row(external_id):
                continue

            slug = slugify(external_id)
            hotel = db.scalar(select(Hotel).where(Hotel.slug == slug))
            if hotel is None:
                hotel = Hotel(slug=slug, name=name)
                db.add(hotel)
                db.flush()
                hotels_created += 1
            else:
                hotels_updated += 1

            hotel.name = name
            hotel.telegram_label = name
            hotel.is_active = True
            hotels_by_external_id[external_id] = hotel
            imported_hotel_ids.add(hotel.id)

        for record in row_dicts(menu_sheet, MENU_REQUIRED_COLUMNS):
            external_id = normalized_text(record.get("Hotel ID"))
            item_name = normalized_text(record.get("Item Name"))
            if not external_id or not item_name or is_summary_row(external_id):
                continue

            hotel = hotels_by_external_id.get(external_id)
            if hotel is None:
                hotel_name = normalized_text(record.get("Hotel Name")) or external_id
                hotel = Hotel(slug=slugify(external_id), name=hotel_name, telegram_label=hotel_name, is_active=True)
                db.add(hotel)
                db.flush()
                hotels_by_external_id[external_id] = hotel
                imported_hotel_ids.add(hotel.id)
                hotels_created += 1

            menu_item = db.scalar(
                select(MenuItem).where(MenuItem.hotel_id == hotel.id, MenuItem.name == item_name)
            )
            if menu_item is None:
                menu_item = MenuItem(hotel_id=hotel.id, name=item_name)
                db.add(menu_item)
                menu_items_created += 1
            else:
                menu_items_updated += 1

            menu_item.description = normalized_text(record.get("Description")) or None
            menu_item.price = decimal_price(record.get("Price (₹)"))
            menu_item.is_available = yes_no(record.get("Available"), default=True)
            imported_menu_keys.add((hotel.id, item_name))

        if deactivate_missing and imported_hotel_ids:
            existing_items = db.scalars(select(MenuItem).where(MenuItem.hotel_id.in_(imported_hotel_ids))).all()
            for menu_item in existing_items:
                if (menu_item.hotel_id, menu_item.name) not in imported_menu_keys and menu_item.is_available:
                    menu_item.is_available = False
                    menu_items_deactivated += 1

        db.commit()
        return ImportResult(
            hotels_created=hotels_created,
            hotels_updated=hotels_updated,
            menu_items_created=menu_items_created,
            menu_items_updated=menu_items_updated,
            menu_items_deactivated=menu_items_deactivated,
        )
    except Exception:
        db.rollback()
        raise
    finally:
        workbook.close()
        if owns_session:
            db.close()


def default_excel_path() -> Path:
    configured = os.getenv("EXCEL_DATABASE_PATH")
    if configured:
        return Path(configured)
    return Path.home() / "OneDrive" / "Desktop" / "MultiHotel_FoodDB.xlsx"


def main() -> None:
    parser = argparse.ArgumentParser(description="Import hotels and menu items from MultiHotel_FoodDB.xlsx.")
    parser.add_argument("--path", type=Path, default=default_excel_path(), help="Path to the Excel workbook.")
    parser.add_argument(
        "--deactivate-missing",
        action="store_true",
        help="Mark menu items as unavailable when they are absent from the workbook.",
    )
    parser.add_argument("--create-tables", action="store_true", help="Create database tables before importing.")
    parser.add_argument("--dry-run", action="store_true", help="Validate the workbook without writing to the database.")
    args = parser.parse_args()

    if args.dry_run:
        summary = inspect_excel(args.path)
        print(
            "Excel dry run complete. "
            f"Hotels: {summary.hotels}, menu items: {summary.menu_items}, "
            f"hotel IDs in menu: {summary.menu_hotel_ids}."
        )
        return

    try:
        if args.create_tables:
            Base.metadata.create_all(bind=engine)
        result = import_excel(args.path, deactivate_missing=args.deactivate_missing)
    except OperationalError as exc:
        print(
            "Could not connect to PostgreSQL. Start the database first, then run this command again.\n\n"
            "From the project root:\n"
            "  docker compose up -d postgres\n"
            "  cd backend\n"
            "  ..\\.venv\\Scripts\\python.exe -m alembic -c alembic.ini upgrade head\n\n"
            f"Database error: {exc.orig}",
            file=sys.stderr,
        )
        raise SystemExit(1) from exc

    print(
        "Excel import complete. "
        f"Hotels created: {result.hotels_created}, hotels updated: {result.hotels_updated}, "
        f"menu items created: {result.menu_items_created}, menu items updated: {result.menu_items_updated}, "
        f"menu items deactivated: {result.menu_items_deactivated}."
    )


if __name__ == "__main__":
    main()
